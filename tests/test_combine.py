"""Tests for Markdown-to-PDF combining with Table of Contents (TOC)."""

from __future__ import annotations

from pathlib import Path

from PySide6.QtCore import QModelIndex
from PySide6.QtGui import QImage
from PySide6.QtPdf import QPdfBookmarkModel, QPdfDocument, QPdfLinkModel

from mdconv.combine import (
    CombineOptions,
    clean_title_from_stem,
    combine_markdown_to_pdf,
    find_markdown_files,
    natural_sort_key,
    prepare_combined_markdown,
)


def test_natural_sort_key():
    files = [
        Path("ch10_advanced.md"),
        Path("ch1_intro.md"),
        Path("ch2_getting_started.md"),
        Path("ch20_appendix.md"),
    ]
    sorted_files = sorted(files, key=natural_sort_key)
    assert [f.name for f in sorted_files] == [
        "ch1_intro.md",
        "ch2_getting_started.md",
        "ch10_advanced.md",
        "ch20_appendix.md",
    ]


def test_clean_title_from_stem():
    assert clean_title_from_stem("01_intro_guide") == "Intro Guide"
    assert clean_title_from_stem("02-system-architecture") == "System Architecture"
    assert clean_title_from_stem("overview") == "Overview"
    assert clean_title_from_stem("10. FAQ") == "FAQ"


def test_find_markdown_files(tmp_path: Path):
    # Setup nested folder structure
    (tmp_path / "01_intro.md").write_text("# Intro", encoding="utf-8")
    sub1 = tmp_path / "part1"
    sub1.mkdir()
    (sub1 / "01_concept.md").write_text("# Concept", encoding="utf-8")
    (sub1 / "02_detail.markdown").write_text("# Detail", encoding="utf-8")
    (sub1 / "ignored.txt").write_text("Ignore", encoding="utf-8")

    sub2 = tmp_path / "part2"
    sub2.mkdir()
    (sub2 / "summary.md").write_text("# Summary", encoding="utf-8")

    found = find_markdown_files([tmp_path], recursive=True)
    names = [f.name for f in found]
    assert names == ["01_intro.md", "01_concept.md", "02_detail.markdown", "summary.md"]


def test_prepare_combined_markdown(tmp_path: Path):
    f1 = tmp_path / "01_intro.md"
    f1.write_text(
        "---\ntitle: 'Spezieller Intro Titel'\nauthor: Tester\n---\n\nErster Text ohne Heading.",
        encoding="utf-8",
    )
    f2 = tmp_path / "02_chapter.md"
    f2.write_text("# Zweites Kapitel\n\nInhalt Kapitel 2.", encoding="utf-8")

    opt = CombineOptions(title="Gesamtwerk")
    combined = prepare_combined_markdown([f1, f2], opt, root_dir=tmp_path)

    # Verify that f1 got an H1 heading with its frontmatter title
    assert "# Spezieller Intro Titel" in combined
    assert "Erster Text ohne Heading." in combined
    assert "# Zweites Kapitel" in combined
    assert "---" in combined  # Main doc header


def test_combine_markdown_to_pdf_with_toc_and_bookmarks(tmp_path: Path):
    # Create structured files
    docs_dir = tmp_path / "docs"
    docs_dir.mkdir()
    sub_dir = docs_dir / "architecture"
    sub_dir.mkdir()

    # Image in subfolder
    img_path = sub_dir / "diagram.png"
    q_img = QImage(20, 20, QImage.Format.Format_RGB32)
    q_img.fill(0x00FF00)
    q_img.save(str(img_path))

    f1 = docs_dir / "01_overview.md"
    f1.write_text("# Gesamtübersicht\n\nEinleitung in das Handbuch.", encoding="utf-8")

    f2 = sub_dir / "01_components.md"
    f2.write_text(
        "# Architekturkomponenten\n\nHier ein Schaubild:\n\n![Architektur](diagram.png)\n\n## Backend Details\n\nDetails zum Backend.",
        encoding="utf-8",
    )

    f3 = docs_dir / "03_conclusion.md"
    f3.write_text(
        "Hier steht Text ohne H1-Überschrift.\n\nUnd noch ein Absatz.",
        encoding="utf-8",
    )

    out_pdf = tmp_path / "handbuch.pdf"
    files = find_markdown_files([docs_dir], recursive=True)

    options = CombineOptions(
        title="Softwarehandbuch",
        author="Data-AL Team",
        toc=True,
        toc_depth=3,
        pagebreaks=True,
        number_sections=True,
    )

    result_pdf = combine_markdown_to_pdf(files, out_pdf, options)
    assert result_pdf.exists()
    assert result_pdf.stat().st_size > 1000

    # Validate PDF structure via PySide6 PDF
    doc = QPdfDocument()
    doc.load(str(result_pdf))
    page_count = doc.pageCount()
    assert page_count >= 4  # Cover, TOC, chapters...

    # Check bookmarks (document outline)
    model = QPdfBookmarkModel()
    model.setDocument(doc)
    assert model.rowCount() >= 3

    # Check clickable links in Table of Contents
    # The TOC is on page 1 (after cover page 0)
    link_model = QPdfLinkModel()
    link_model.setDocument(doc)
    link_model.setPage(1)
    toc_links_count = link_model.rowCount(QModelIndex())
    assert toc_links_count >= 3  # Links to the chapters

    doc.close()


def test_sanitize_typst_content(tmp_path: Path):
    from mdconv.combine import sanitize_typst_content

    existing_img = tmp_path / "valid.png"
    existing_img.touch()

    typ_input = (
        '#box(image("https://example.com/logo.png"))\n'
        '#figure(image("missing.png"), caption: [Missing])\n'
        '#figure(image("valid.png"), caption: [Valid])\n'
    )

    sanitized = sanitize_typst_content(typ_input, common_root=tmp_path)
    assert 'link("https://example.com/logo.png")' in sanitized
    assert "Bild nicht gefunden: missing.png" in sanitized
    assert 'image("valid.png")' in sanitized


def test_combine_with_remote_and_missing_images(tmp_path: Path):
    """Ensure remote images and missing local images do not crash PDF compilation."""
    f1 = tmp_path / "01_readme.md"
    f1.write_text(
        "# Projekt mit Bildern\n\n"
        "Ein lokales Bild das fehlt:\n\n"
        "![Missing local](images/does_not_exist.png)\n\n"
        "Ein Online-Badge / Bild:\n\n"
        "![Remote Badge](https://user-images.githubusercontent.com/5280692/170442888-f64c2a1b-c99d-46fa-8e76-210706b484b0.png)\n\n"
        "Ein nicht erreichbares Online-Bild:\n\n"
        "![Unreachable](https://httpbin.org/status/404)\n",
        encoding="utf-8",
    )

    out_pdf = tmp_path / "combined_resilient.pdf"
    res = combine_markdown_to_pdf([f1], out_pdf, CombineOptions(title="Resilienz-Test"))
    assert res.exists()
    assert res.stat().st_size > 500


def test_combine_multiple_file_types(tmp_path: Path):
    """Test combining various formats (MD, TXT, CSV, XLSX, JSON, PY, PNG) into a single PDF."""
    from openpyxl import Workbook

    from mdconv.combine import find_combinable_files

    docs = tmp_path / "mixed_docs"
    docs.mkdir()

    # 1. Markdown
    (docs / "01_intro.md").write_text(
        "# Projektübersicht\n\nDas ist die Einleitung.", encoding="utf-8"
    )

    # 2. Text file
    (docs / "02_notes.txt").write_text(
        "Wichtige Notizen zum Projektverlauf.", encoding="utf-8"
    )

    # 3. CSV
    (docs / "03_data.csv").write_text(
        "ID,Name,Status\n1,Alpha,Aktiv\n2,Beta,Geplant\n", encoding="utf-8"
    )

    # 4. XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "Kennzahlen"
    ws.append(["Metrik", "Wert"])
    ws.append(["Umsatz", 150000])
    wb.save(docs / "04_metrics.xlsx")

    # 5. JSON code/config
    (docs / "05_config.json").write_text(
        '{\n  "version": "1.0",\n  "enabled": true\n}\n', encoding="utf-8"
    )

    # 6. Python script
    (docs / "06_script.py").write_text(
        'def run():\n    print("System aktiv")\n', encoding="utf-8"
    )

    # 7. Image
    img_path = docs / "07_schema.png"
    q_img = QImage(30, 30, QImage.Format.Format_RGB32)
    q_img.fill(0x0000FF)
    q_img.save(str(img_path))

    # Discover combinable files
    files = find_combinable_files([docs], recursive=True)
    assert len(files) == 7

    out_pdf = tmp_path / "gesamt_dokumentation.pdf"
    opt = CombineOptions(title="Gesamthandbuch", toc=True, toc_depth=3)
    res = combine_markdown_to_pdf(files, out_pdf, opt)

    assert res.exists()
    assert res.stat().st_size > 1000

    doc = QPdfDocument()
    doc.load(str(res))
    assert doc.pageCount() >= 5

    # Check outline bookmarks
    model = QPdfBookmarkModel()
    model.setDocument(doc)
    assert model.rowCount() >= 6
    doc.close()


def test_combine_with_scoped_packages_and_anchor_links(tmp_path: Path):
    """Ensure npm scoped package mentions (@foo/bar) and GitHub/GitLab anchor links compile smoothly."""
    f1 = tmp_path / "01_gitlab_readme.md"
    f1.write_text(
        "# @data-al/design-system\n\n"
        "Paket installieren:\n\n"
        "- npm install @data-al/design-system\n"
        "- Plugin: [Vite Plugin](https://github.com) [@vitejs/plugin-react]\n\n"
        "Inhaltsverzeichnis:\n"
        "- [Installation & Setup](#installation--setup)\n"
        "- [1. Enqueue tables](#1-enqueue-tables)\n"
        "- [Accounts & workspaces](#accounts--workspaces)\n"
        "- [Dead link to nowhere](#unknown-heading-nowhere)\n\n"
        "# Installation & Setup\n\n"
        "Hier steht die Installation.\n\n"
        "# 1. Enqueue tables\n\n"
        "Hier stehen die Queue-Details.\n\n"
        "# Accounts & workspaces\n\n"
        "Hier stehen die Konten.\n",
        encoding="utf-8",
    )

    out_pdf = tmp_path / "gitlab_test.pdf"
    res = combine_markdown_to_pdf([f1], out_pdf, CombineOptions(title="GitLab Test"))
    assert res.exists()
    assert res.stat().st_size > 1000


def test_combine_with_html_document(tmp_path: Path):
    """Ensure HTML files are parsed and combined via Pandoc."""
    html_file = tmp_path / "page.html"
    html_file.write_text(
        "<h1>Web Dokument</h1><p>Ein Absatz mit <strong>fettem Text</strong>.</p>",
        encoding="utf-8",
    )
    out_pdf = tmp_path / "html_test.pdf"
    res = combine_markdown_to_pdf(
        [html_file], out_pdf, CombineOptions(title="HTML Test")
    )
    assert res.exists()
    assert res.stat().st_size > 1000
