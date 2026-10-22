import importlib.util
import shutil

import pytest
from openpyxl import load_workbook

from mdconv.cli import main as cli
from mdconv.engine import ConversionError, Options, collect_inputs, convert

SAMPLE = """# Bericht

Text mit **fett** und *kursiv*.

## Umsatz

| Monat | Betrag |
|---|---|
| Jan | 1200 |
| Feb | 980.5 |

## Personal

| Name | Team |
|---|---|
| Anna | Support |
"""

needs_pandoc = pytest.mark.skipif(
    not (shutil.which("pandoc") or importlib.util.find_spec("pypandoc")),
    reason="pandoc nicht verfügbar",
)


@pytest.fixture
def md(tmp_path):
    p = tmp_path / "bericht.md"
    p.write_text(SAMPLE, encoding="utf-8")
    return p


def test_md_xlsx_roundtrip(md, tmp_path):
    [xlsx] = convert(md, "xlsx", tmp_path / "out")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Umsatz", "Personal"]
    assert [c.value for c in wb["Umsatz"][3]] == ["Feb", 980.5]

    [back] = convert(xlsx, "md", tmp_path / "back")
    text = back.read_text(encoding="utf-8")
    assert "## Umsatz" in text and "| Feb | 980.5 |" in text


def test_md_to_csv_writes_one_file_per_table(md, tmp_path):
    outs = convert(md, "csv", tmp_path / "out", Options(csv_delimiter=";"))
    assert [o.name for o in outs] == ["bericht_1.csv", "bericht_2.csv"]
    assert outs[0].read_text(encoding="utf-8-sig").splitlines()[1] == "Jan;1200"


def test_csv_to_md_sniffs_semicolon(tmp_path):
    src = tmp_path / "daten.csv"
    src.write_text("Name;Ort\nAnna;Stuttgart\n", encoding="utf-8")
    [out] = convert(src, "md")
    assert "| Anna | Stuttgart |" in out.read_text(encoding="utf-8")


def test_existing_target_is_protected(md, tmp_path):
    convert(md, "xlsx", tmp_path)
    with pytest.raises(ConversionError, match="existiert bereits"):
        convert(md, "xlsx", tmp_path)
    convert(md, "xlsx", tmp_path, Options(overwrite=True))


def test_unsupported_and_tableless(tmp_path):
    txt = tmp_path / "a.txt"
    txt.write_text("x")
    with pytest.raises(ConversionError, match="nicht unterstützt"):
        convert(txt, "md")
    plain = tmp_path / "plain.md"
    plain.write_text("# nur Text\n")
    with pytest.raises(ConversionError, match="Keine Markdown-Tabelle"):
        convert(plain, "xlsx")


def test_collect_inputs_filters_by_target(tmp_path):
    (tmp_path / "a.md").write_text("x")
    (tmp_path / "b.docx").write_bytes(b"")
    sub = tmp_path / "sub"
    sub.mkdir()
    (sub / "c.md").write_text("x")
    assert [p.name for p in collect_inputs([tmp_path], "docx")] == ["a.md"]
    assert [p.name for p in collect_inputs([tmp_path], "docx", recursive=True)] == [
        "a.md",
        "c.md",
    ]
    assert [p.name for p in collect_inputs([tmp_path], "md")] == ["b.docx"]


@needs_pandoc
def test_md_docx_roundtrip_keeps_structure(md, tmp_path):
    [docx] = convert(md, "docx", tmp_path / "out")
    [back] = convert(docx, "md", tmp_path / "back")
    text = back.read_text(encoding="utf-8")
    assert "# Bericht" in text and "**fett**" in text
    assert "| Anna | Support |" in text


@needs_pandoc
def test_md_html_roundtrip(md, tmp_path):
    [html] = convert(md, "html", tmp_path / "out")
    assert "<table" in html.read_text(encoding="utf-8")
    [back] = convert(html, "md", tmp_path / "back")
    assert "Support" in back.read_text(encoding="utf-8")


@pytest.mark.skipif(
    importlib.util.find_spec("docling") is not None, reason="docling installiert"
)
def test_pdf_import_without_docling_gives_clear_message(tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    with pytest.raises(ConversionError, match="PDF-Import"):
        convert(pdf, "md")


def test_cli_batch_reports_failures(md, tmp_path, capsys):
    (tmp_path / "plain.md").write_text("# ohne Tabelle\n")
    code = cli([str(tmp_path), "--to", "xlsx", "-o", str(tmp_path / "out")])
    out = capsys.readouterr()
    assert code == 1
    assert "1 von 2 Dateien konvertiert" in out.out
    assert "plain.md" in out.err
