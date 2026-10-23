"""Combine multiple Markdown files into a single structured PDF with a clickable Table of Contents (TOC)."""

from __future__ import annotations

import base64
import csv
import hashlib
import os
import re
import shutil
import tempfile
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING

from mdconv.tools import ToolMissingError, find_tool, run

if TYPE_CHECKING:
    from collections.abc import Iterable


COMBINABLE_EXTENSIONS: set[str] = {
    # Markdown
    ".md",
    ".markdown",
    # Rich Office & Web Documents
    ".docx",
    ".odt",
    ".html",
    ".htm",
    ".rst",
    # Tables & Data
    ".xlsx",
    ".csv",
    ".tsv",
    # PDF
    ".pdf",
    # Plain text & Logs
    ".txt",
    ".log",
    # Source Code & Config
    ".py",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".sh",
    ".bash",
    ".bat",
    ".cmd",
    ".ps1",
    ".sql",
    ".css",
    ".scss",
    ".java",
    ".cpp",
    ".c",
    ".h",
    ".hpp",
    ".rs",
    ".go",
    ".php",
    ".rb",
    ".r",
    # Standalone Images
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".svg",
    ".webp",
}

IGNORE_DIR_NAMES: set[str] = {
    ".git",
    ".svn",
    ".hg",
    ".venv",
    "venv",
    "env",
    "__pycache__",
    ".idea",
    ".vscode",
    "node_modules",
}

CODE_LANG_MAP: dict[str, str] = {
    ".py": "python",
    ".js": "javascript",
    ".ts": "typescript",
    ".jsx": "jsx",
    ".tsx": "tsx",
    ".json": "json",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".xml": "xml",
    ".toml": "toml",
    ".ini": "ini",
    ".cfg": "ini",
    ".conf": "ini",
    ".sh": "bash",
    ".bash": "bash",
    ".bat": "bat",
    ".cmd": "bat",
    ".ps1": "powershell",
    ".sql": "sql",
    ".css": "css",
    ".scss": "scss",
    ".java": "java",
    ".cpp": "cpp",
    ".c": "c",
    ".h": "c",
    ".hpp": "cpp",
    ".rs": "rust",
    ".go": "go",
    ".php": "php",
    ".rb": "ruby",
    ".r": "r",
}


@dataclass
class CombineOptions:
    """Options for merging multiple Markdown and document files into a single PDF document."""

    title: str = "Dokumentation"
    author: str | None = None
    date: str | None = None
    toc: bool = True
    toc_depth: int = 3
    pagebreaks: bool = True
    number_sections: bool = False
    paper_size: str = "a4"
    overwrite: bool = False
    auto_headings: bool = True


def is_combinable_file(path: Path) -> bool:
    """Checks whether a file can be incorporated into the combined PDF."""
    if not path.is_file():
        return False
    name = path.name
    if name.startswith((".", "~$")) or name.endswith("~"):
        return False
    for parent in path.parents:
        if parent.name in IGNORE_DIR_NAMES:
            return False
    return path.suffix.lower() in COMBINABLE_EXTENSIONS


def natural_sort_key(path: Path) -> list[int | str]:
    """Provides a natural alphanumeric sort key (e.g. 'part2' comes before 'part10')."""
    parts: list[int | str] = []
    # Include relative path string for consistent folder + filename ordering
    for component in re.split(r"(\d+)", path.as_posix()):
        if not component:
            continue
        parts.append(int(component) if component.isdigit() else component.lower())
    return parts


def find_markdown_files(paths: Iterable[Path], recursive: bool = True) -> list[Path]:
    """Collects and naturally sorts all Markdown files from the given paths and folders."""
    found: list[Path] = []
    for raw in paths:
        p = Path(raw).resolve()
        if p.is_file() and p.suffix.lower() in {".md", ".markdown"}:
            found.append(p)
        elif p.is_dir():
            pattern = "**/*" if recursive else "*"
            for item in p.glob(pattern):
                if item.is_file() and item.suffix.lower() in {".md", ".markdown"}:
                    found.append(item.resolve())

    unique = list(dict.fromkeys(found))
    unique.sort(key=natural_sort_key)
    return unique


def find_combinable_files(paths: Iterable[Path], recursive: bool = True) -> list[Path]:
    """Collects and naturally sorts all combinable files from the given paths and folders."""
    found: list[Path] = []
    for raw in paths:
        p = Path(raw).resolve()
        if is_combinable_file(p):
            found.append(p)
        elif p.is_dir():
            pattern = "**/*" if recursive else "*"
            for item in p.glob(pattern):
                if is_combinable_file(item):
                    found.append(item.resolve())

    unique = list(dict.fromkeys(found))
    unique.sort(key=natural_sort_key)
    return unique


def _get_pdf_page_count(path: Path) -> int:
    """Returns number of pages in a PDF document using PySide6 or regex fallback."""
    try:
        import gc

        from PySide6.QtPdf import QPdfDocument

        doc = QPdfDocument()
        doc.load(str(path))
        count = doc.pageCount()
        doc.close()
        del doc
        gc.collect()
        if count > 0:
            return count
    except (ImportError, RuntimeError, OSError, ValueError):
        count = 0

    # Fallback: scan binary for /Type /Page
    try:
        data = path.read_bytes()
        found_pages = len(re.findall(rb"/Type\s*/Page\b(?!s)", data))
        return max(found_pages, 1)
    except (OSError, ValueError):
        return 1


def file_to_markdown(file_path: Path, root_dir: Path, assets_dir: Path) -> str:
    """Converts any supported file into Markdown representation for combining."""
    ext = file_path.suffix.lower()

    # 1. Native Markdown
    if ext in {".md", ".markdown"}:
        return file_path.read_text(encoding="utf-8-sig", errors="replace")

    # 2. Rich Office & Web Documents via Pandoc (.docx, .odt, .html, .htm, .rst)
    if ext in {".docx", ".odt", ".html", ".htm", ".rst"}:
        try:
            pandoc = find_tool("pandoc")
            fmt_map = {
                ".docx": "docx",
                ".odt": "odt",
                ".html": "html",
                ".htm": "html",
                ".rst": "rst",
            }
            fmt = fmt_map[ext]
            args = [pandoc, str(file_path), "-f", fmt, "-t", "gfm", "--wrap=none"]
            if ext in {".docx", ".odt"}:
                media_sub = assets_dir / f"media_{file_path.stem}"
                args.append(f"--extract-media={media_sub}")
            proc = run(args, cwd=root_dir)
            return proc.stdout if proc and proc.stdout else ""
        except (
            OSError,
            RuntimeError,
            ValueError,
            ToolMissingError,
            AttributeError,
        ) as exc:
            return f"*[Fehler beim Laden von {file_path.name}: {exc}]*"

    # 3. Excel Spreadsheets (.xlsx)
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook

            from mdconv import tables

            wb = load_workbook(file_path, read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                rows = tables.trim([list(r) for r in ws.iter_rows(values_only=True)])
                if rows:
                    parts.append(f"## {ws.title}\n\n{tables.to_markdown(rows)}")
            wb.close()
            return (
                "\n\n".join(parts) if parts else "*[Arbeitsmappe enthält keine Daten]*"
            )
        except (OSError, KeyError, ValueError) as exc:
            return f"*[Fehler beim Laden der Excel-Tabelle {file_path.name}: {exc}]*"

    # 4. CSV & TSV Data (.csv, .tsv)
    if ext in {".csv", ".tsv"}:
        try:
            from mdconv import tables

            text = file_path.read_text(encoding="utf-8-sig", errors="replace")
            try:
                delimiter = "\t" if ext == ".tsv" else None
                if delimiter:
                    rows = tables.trim(
                        list(csv.reader(text.splitlines(), delimiter=delimiter))
                    )
                else:
                    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
                    rows = tables.trim(list(csv.reader(text.splitlines(), dialect)))
            except csv.Error:
                rows = tables.trim(list(csv.reader(text.splitlines())))
            return tables.to_markdown(rows) if rows else "*[Tabelle ist leer]*"
        except (OSError, ValueError) as exc:
            return f"*[Fehler beim Laden der Datei {file_path.name}: {exc}]*"

    # 5. PDF Documents (.pdf)
    if ext == ".pdf":
        try:
            import importlib

            docling_dc = importlib.import_module("docling.document_converter")
            converter = docling_dc.DocumentConverter()
            result = converter.convert(file_path)
            return result.document.export_to_markdown()
        except (ImportError, OSError, ValueError):
            num_pages = _get_pdf_page_count(file_path)
            if num_pages > 0:
                try:
                    rel_pdf = file_path.relative_to(root_dir).as_posix()
                except ValueError:
                    assets_dir.mkdir(parents=True, exist_ok=True)
                    copied = assets_dir / f"pdf_{file_path.name}"
                    if not copied.exists():
                        shutil.copy2(file_path, copied)
                    rel_pdf = copied.relative_to(root_dir).as_posix()

                typst_cmds: list[str] = []
                for p in range(1, num_pages + 1):
                    typst_cmds.append(f'#image("{rel_pdf}", page: {p})')
                    if p < num_pages:
                        typst_cmds.append("#pagebreak()")
                return "```{=typst}\n" + "\n".join(typst_cmds) + "\n```"
            return f"*[PDF-Datei {file_path.name} konnte nicht eingebettet werden]*"

    # 6. Standalone Images (.png, .jpg, .jpeg, .gif, .svg, .webp)
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp"}:
        try:
            rel_img = file_path.relative_to(root_dir).as_posix()
        except ValueError:
            assets_dir.mkdir(parents=True, exist_ok=True)
            copied = assets_dir / f"img_{file_path.name}"
            if not copied.exists():
                shutil.copy2(file_path, copied)
            rel_img = copied.relative_to(root_dir).as_posix()
        title = clean_title_from_stem(file_path.stem)
        return f"![{title}]({rel_img})"

    # 7. Code & Config with syntax highlighting
    if ext in CODE_LANG_MAP:
        lang = CODE_LANG_MAP[ext]
        text = file_path.read_text(encoding="utf-8-sig", errors="replace")
        return f"```{lang}\n{text}\n```"

    # 8. Plain text & Logs (.txt, .log)
    return file_path.read_text(encoding="utf-8-sig", errors="replace")


def clean_title_from_stem(stem: str) -> str:
    """Derives a human-readable title from a file stem (e.g. '01_intro_guide' -> 'Intro Guide')."""
    # Remove leading number prefixes like "01_", "1-", "01. "
    cleaned = re.sub(r"^\d+[\s_.-]+", "", stem)
    # Replace underscores and hyphens with spaces
    cleaned = re.sub(r"[_-]+", " ", cleaned).strip()
    if not cleaned:
        cleaned = stem
    # Capitalize words if all lowercase
    if cleaned.islower():
        cleaned = cleaned.title()
    return cleaned


_FRONTMATTER_PATTERN = re.compile(r"^---\s*\n(.*?)\n---\s*\n?", re.DOTALL)
_H1_PATTERN = re.compile(r"^\s*#\s+(.+)$", re.MULTILINE)
_IMG_PATTERN = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def _extract_title_from_frontmatter(yaml_text: str) -> str | None:
    """Extracts 'title: ...' from raw YAML frontmatter text."""
    for line in yaml_text.splitlines():
        line = line.strip()
        if line.startswith("title:"):
            raw_val = line[6:].strip().strip("\"'")
            if raw_val:
                return raw_val
    return None


def _detect_image_extension(data: bytes, url_or_name: str) -> str | None:
    """Detects supported image extension (.png, .jpg, .gif, .svg) from magic bytes or URL."""
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return ".png"
    if data.startswith(b"\xff\xd8\xff"):
        return ".jpg"
    if data.startswith((b"GIF87a", b"GIF89a")):
        return ".gif"

    # SVG check: XML or SVG tag in header
    snippet = data[:2048].lstrip()
    if snippet.startswith(b"<svg") or (b"<?xml" in snippet and b"<svg" in snippet):
        return ".svg"

    # Fallback to URL / filename extension
    clean = url_or_name.split("?")[0].split("#")[0].lower()
    for ext in (".png", ".jpg", ".jpeg", ".svg", ".gif"):
        if clean.endswith(ext):
            return ".jpg" if ext == ".jpeg" else ext

    return None


def _download_and_cache_image(
    url: str,
    assets_dir: Path,
    cache: dict[str, Path | None],
) -> Path | None:
    """Downloads a remote image to assets_dir if valid and supported by Typst."""
    if url in cache:
        return cache[url]

    if not url.startswith(("http://", "https://")):
        cache[url] = None
        return None

    try:
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
                ),
                "Accept": "image/*,*/*;q=0.8",
            },
        )
        with urllib.request.urlopen(req, timeout=5.0) as resp:
            status = getattr(resp, "status", 200)
            if status != 200:
                cache[url] = None
                return None
            data = resp.read(25 * 1024 * 1024)  # 25 MB limit

        ext = _detect_image_extension(data, url)
        if not ext:
            cache[url] = None
            return None

        assets_dir.mkdir(parents=True, exist_ok=True)
        url_hash = hashlib.sha256(url.encode("utf-8")).hexdigest()[:16]
        dest_file = assets_dir / f"img_{url_hash}{ext}"
        dest_file.write_bytes(data)
        cache[url] = dest_file
        return dest_file
    except (urllib.error.URLError, TimeoutError, OSError, ValueError):
        cache[url] = None
        return None


def _save_data_uri_image(
    data_uri: str,
    assets_dir: Path,
    counter: list[int],
) -> Path | None:
    """Decodes a base64 data URI image and saves it to assets_dir."""
    match = re.match(r"^data:image/([a-zA-Z0-9+.-]+);base64,(.+)$", data_uri, re.DOTALL)
    if not match:
        return None
    raw_format, b64_data = match.groups()
    ext_map = {
        "png": ".png",
        "jpeg": ".jpg",
        "jpg": ".jpg",
        "gif": ".gif",
        "svg+xml": ".svg",
    }
    ext = ext_map.get(raw_format.lower())
    if not ext:
        return None
    try:
        data = base64.b64decode(b64_data)
        assets_dir.mkdir(parents=True, exist_ok=True)
        counter[0] += 1
        dest_file = assets_dir / f"data_img_{counter[0]}{ext}"
        dest_file.write_bytes(data)
        return dest_file
    except (ValueError, OSError):
        return None


def sanitize_typst_content(typ_content: str, common_root: Path) -> str:
    """Sanitizes Typst file content to prevent compilation errors from remote/missing images,

    missing labels in internal links, and citations without bibliography.
    """
    # 1. Sanitize image(...) calls
    pattern = re.compile(r'image\(\s*"([^"]+)"(?P<extra>[^)]*)\)')

    def _replace_typ_image(match: re.Match) -> str:
        src = match.group(1)
        if src.startswith(("http://", "https://")):
            return f'link("{src}")[[Bild: {src}]]'
        try:
            norm_path = Path(src.replace("\\", "/"))
            resolved = (common_root / norm_path).resolve()
            resolved.relative_to(common_root.resolve())
            if not resolved.is_file():
                return f"[(Bild nicht gefunden: {src})]"
        except (ValueError, OSError):
            return f"[(Bild nicht gefunden: {src})]"
        return match.group(0)

    typ_content = pattern.sub(_replace_typ_image, typ_content)

    # 2. Sanitize #cite(...) calls when no bibliography is defined in document
    cite_pattern = re.compile(r'#?cite\(\s*label\("([^"]+)"\)[^)]*\)')
    typ_content = cite_pattern.sub(r"[\@\1]", typ_content)

    # 3. Resolve internal links and normalize labels (e.g. accounts--workspaces -> accounts-workspaces)
    defined_labels = set(
        re.findall(r"(?<!#link\()<([a-zA-Z0-9_\-.:]+)>(?!\s*\[)", typ_content)
    )

    def normalize_label(lbl: str) -> str:
        s = re.sub(r"-+", "-", lbl).strip("-")
        s_no_num = re.sub(r"^\d+-*", "", s)
        return s_no_num.lower()

    norm_to_actual = {normalize_label(lbl): lbl for lbl in defined_labels}
    referenced = set(re.findall(r"#?link\(\s*<([a-zA-Z0-9_\-.:]+)>\s*\)", typ_content))

    missing_labels: list[str] = []
    for ref in referenced:
        if ref in defined_labels:
            continue
        norm = normalize_label(ref)
        if norm in norm_to_actual:
            actual = norm_to_actual[norm]
            typ_content = typ_content.replace(f"<{ref}>", f"<{actual}>")
        else:
            missing_labels.append(ref)

    # 4. Inject dummy targets for any remaining missing labels so Typst never fails
    if missing_labels:
        dummy_defs = "\n".join(f"#place(hide([])) <{lbl}>" for lbl in missing_labels)
        typ_content = dummy_defs + "\n\n" + typ_content

    return typ_content


def prepare_combined_markdown(
    files: list[Path],
    options: CombineOptions,
    root_dir: Path,
    assets_dir: Path | None = None,
) -> str:
    """Merges Markdown files into a unified document with adjusted image links and chapter headings."""
    if assets_dir is None:
        assets_dir = root_dir / ".mdconv_assets"

    url_cache: dict[str, Path | None] = {}
    data_counter: list[int] = [0]
    sections: list[str] = []

    # Document Header with metadata for Pandoc
    meta_lines = ["---", f'title: "{options.title}"', 'lang: "de"']
    if options.author:
        meta_lines.append(f'author: "{options.author}"')
    if options.date:
        meta_lines.append(f'date: "{options.date}"')
    meta_lines.append("---\n")
    sections.append("\n".join(meta_lines))

    for file_path in files:
        raw_text = file_to_markdown(file_path, root_dir=root_dir, assets_dir=assets_dir)

        # Extract & strip frontmatter
        extracted_title: str | None = None
        match = _FRONTMATTER_PATTERN.match(raw_text)
        if match:
            extracted_title = _extract_title_from_frontmatter(match.group(1))
            body = raw_text[match.end() :]
        else:
            body = raw_text

        body = body.strip()

        # Check if the file already contains an H1 heading
        has_h1 = bool(_H1_PATTERN.search(body))

        if not has_h1 and options.auto_headings:
            # Generate an H1 heading so this file appears in the Table of Contents
            heading_title = extracted_title or clean_title_from_stem(file_path.stem)
            body = f"# {heading_title}\n\n{body}"

        # Adjust relative and remote image paths
        def _replace_img(match_obj: re.Match, current_file: Path = file_path) -> str:
            alt = match_obj.group(1)
            raw_src = match_obj.group(2).strip()

            title_part = ""
            src = raw_src
            if ' "' in src:
                src, title_part = src.split(' "', 1)
                title_part = f' "{title_part}'
            elif " '" in src:
                src, title_part = src.split(" '", 1)
                title_part = f" '{title_part}"

            src = src.strip()

            # Ignore pure anchors
            if src.startswith("#"):
                return match_obj.group(0)

            # 1. Remote HTTP/HTTPS images: download and cache for Typst
            if src.startswith(("http://", "https://")):
                cached = _download_and_cache_image(src, assets_dir, url_cache)
                if cached:
                    try:
                        rel_path = cached.relative_to(root_dir).as_posix()
                        return f"![{alt}]({rel_path}{title_part})"
                    except ValueError:
                        return f"![{alt}]({cached.as_posix()}{title_part})"
                # Fallback to link if download failed or format is unsupported
                link_text = alt or "Bild"
                return f"[{link_text}]({src})"

            # 2. Data URI images
            if src.startswith("data:image/"):
                saved = _save_data_uri_image(src, assets_dir, data_counter)
                if saved:
                    try:
                        rel_path = saved.relative_to(root_dir).as_posix()
                        return f"![{alt}]({rel_path}{title_part})"
                    except ValueError:
                        return f"![{alt}]({saved.as_posix()}{title_part})"
                return f"*[Bild: {alt or 'Daten-Bild'}]*"

            # 3. Local images
            clean_src = src.split("?")[0].split("#")[0]
            if (len(clean_src) > 1 and clean_src[1] == ":") or clean_src.startswith(
                ("/", "\\")
            ):
                img_file = Path(clean_src).resolve()
            else:
                img_file = (current_file.parent / clean_src).resolve()

            if not img_file.is_file():
                return f"*[Bild nicht gefunden: {alt or src}]*"

            try:
                rel_to_root = img_file.relative_to(root_dir).as_posix()
                return f"![{alt}]({rel_to_root}{title_part})"
            except ValueError:
                # Image is outside root_dir: copy it into assets_dir so Typst sandbox permits it
                try:
                    assets_dir.mkdir(parents=True, exist_ok=True)
                    path_hash = hashlib.sha256(
                        str(img_file).encode("utf-8")
                    ).hexdigest()[:12]
                    copied_img = assets_dir / f"local_{path_hash}_{img_file.name}"
                    if not copied_img.exists():
                        shutil.copy2(img_file, copied_img)
                    rel_to_root = copied_img.relative_to(root_dir).as_posix()
                    return f"![{alt}]({rel_to_root}{title_part})"
                except (OSError, ValueError):
                    return f"![{alt}]({img_file.as_posix()}{title_part})"

        body = _IMG_PATTERN.sub(_replace_img, body)

        sections.append(body)

    return "\n\n".join(sections) + "\n"


def combine_markdown_to_pdf(
    files: list[Path],
    out_file: Path,
    options: CombineOptions | None = None,
) -> Path:
    """Combines a list of Markdown files into a single PDF with a clickable Table of Contents.

    Uses Pandoc and Typst.
    """
    if not files:
        raise ValueError("Keine Markdown-Dateien zum Kombinieren übergeben.")

    opt = options or CombineOptions()
    out_file = Path(out_file).resolve()

    if out_file.exists() and not opt.overwrite:
        raise FileExistsError(
            f"Zieldatei existiert bereits: {out_file.name} (Überschreiben aktivieren)"
        )

    out_file.parent.mkdir(parents=True, exist_ok=True)

    # Determine common root directory for all files
    try:
        common_root = Path(os.path.commonpath([str(f.resolve().parent) for f in files]))
    except (ValueError, OSError):
        common_root = files[0].parent.resolve()

    pandoc = find_tool("pandoc")
    typst = find_tool("typst")

    pid = os.getpid()
    assets_dir = common_root / f".mdconv_assets_{pid}"
    temp_typ = common_root / f".mdconv_combined_{pid}.typ"

    # Clean up legacy unversioned temporary files if present
    legacy_typ = common_root / ".mdconv_combined.typ"
    if legacy_typ.is_file():
        try:
            legacy_typ.unlink()
        except OSError:
            pass
    legacy_assets = common_root / ".mdconv_assets"
    if legacy_assets.is_dir():
        shutil.rmtree(legacy_assets, ignore_errors=True)

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        merged_md = temp_path / "combined.md"

        # Generate unified markdown content
        content = prepare_combined_markdown(
            files, opt, root_dir=common_root, assets_dir=assets_dir
        )
        merged_md.write_text(content, encoding="utf-8")

        # Step 1: Pandoc compiles Markdown to Standalone Typst (.typ)
        pandoc_args = [
            pandoc,
            str(merged_md),
            "-f",
            "markdown-citations",
            "-s",
            "-o",
            str(temp_typ),
            "-M",
            "lang=de",
            "-M",
            f"paper={opt.paper_size}",
        ]

        if opt.toc:
            pandoc_args.extend(["--toc", f"--toc-depth={opt.toc_depth}"])

        if opt.number_sections:
            pandoc_args.append("--number-sections")

        if opt.pagebreaks:
            pandoc_args.extend(
                [
                    "-V",
                    "header-includes=#show heading.where(level: 1): it => pagebreak(weak: true) + it",
                ]
            )

        try:
            run(pandoc_args, cwd=common_root)

            # Sanitize .typ file before Typst compilation
            if temp_typ.is_file():
                typ_text = temp_typ.read_text(encoding="utf-8")
                sanitized_typ = sanitize_typst_content(
                    typ_text, common_root=common_root
                )
                temp_typ.write_text(sanitized_typ, encoding="utf-8")

            # Step 2: Typst compiles .typ to PDF, resolving all images relative to common_root
            typst_args = [
                typst,
                "compile",
                "--root",
                str(common_root),
                str(temp_typ),
                str(out_file),
            ]
            run(typst_args, cwd=common_root)
        finally:
            if temp_typ.exists():
                try:
                    temp_typ.unlink()
                except OSError:
                    pass
            if assets_dir.exists():
                shutil.rmtree(assets_dir, ignore_errors=True)

    return out_file


# Alias for multi-format combining
combine_files_to_pdf = combine_markdown_to_pdf
