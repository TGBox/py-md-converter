"""Konvertierungs-Engine: Registry der Konverter und Batch-Verarbeitung."""

from __future__ import annotations

import csv
import re
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from pathlib import Path

from mdconv import tables
from mdconv.combine import CombineOptions, combine_markdown_to_pdf
from mdconv.tools import ToolMissingError, find_tool, run

__all__ = [
    "CombineOptions",
    "ConversionError",
    "Options",
    "collect_inputs",
    "combine_markdown_to_pdf",
    "convert",
    "format_of",
    "supported_conversions",
]


class ConversionError(Exception):
    pass


@dataclass
class Options:
    reference_doc: Path | None = None  # Word-Vorlage für MD -> DOCX
    csv_delimiter: str = ","
    overwrite: bool = False


EXTENSIONS = {
    ".md": "md",
    ".markdown": "md",
    ".docx": "docx",
    ".html": "html",
    ".htm": "html",
    ".pdf": "pdf",
    ".xlsx": "xlsx",
    ".csv": "csv",
}

Converter = Callable[[Path, Path, Options], list[Path]]
_CONVERTERS: dict[tuple[str, str], Converter] = {}


def _register(src: str, dst: str):
    def deco(fn: Converter) -> Converter:
        _CONVERTERS[(src, dst)] = fn
        return fn

    return deco


def supported_conversions() -> list[tuple[str, str]]:
    return sorted(_CONVERTERS)


def format_of(path: Path) -> str | None:
    return EXTENSIONS.get(path.suffix.lower())


# ---------- Pandoc ----------


def _pandoc(*args: str, cwd: Path | None = None) -> None:
    run([find_tool("pandoc"), *args], cwd=cwd)


@_register("md", "docx")
def _md_to_docx(src: Path, out: Path, opt: Options) -> list[Path]:
    args = [str(src), "-o", str(out), f"--resource-path={src.parent}"]
    if opt.reference_doc:
        args.append(f"--reference-doc={opt.reference_doc}")
    _pandoc(*args)
    return [out]


@_register("md", "html")
def _md_to_html(src: Path, out: Path, opt: Options) -> list[Path]:
    _pandoc(
        str(src),
        "-o",
        str(out),
        "-s",
        "--embed-resources",
        f"--resource-path={src.parent}",
        f"--metadata=pagetitle:{src.stem}",
    )
    return [out]


@_register("md", "pdf")
def _md_to_pdf(src: Path, out: Path, opt: Options) -> list[Path]:
    _pandoc(
        str(src),
        "-o",
        str(out),
        f"--pdf-engine={find_tool('typst')}",
        f"--resource-path={src.parent}",
    )
    return [out]


def _to_markdown(src: Path, out: Path, fmt: str) -> list[Path]:
    # Bilder landen in <name>_media neben der Ausgabe, relativ verlinkt.
    media = f"{out.stem}_media"
    _pandoc(
        str(src),
        "-f",
        fmt,
        "-t",
        "gfm",
        "--wrap=none",
        f"--extract-media={media}",
        "-o",
        out.name,
        cwd=out.parent,
    )
    return [out]


@_register("docx", "md")
def _docx_to_md(src: Path, out: Path, opt: Options) -> list[Path]:
    return _to_markdown(src, out, "docx")


@_register("html", "md")
def _html_to_md(src: Path, out: Path, opt: Options) -> list[Path]:
    return _to_markdown(src, out, "html")


# ---------- Tabellen ----------


def _read_tables(src: Path) -> list[tables.Table]:
    found = tables.parse_tables(src.read_text(encoding="utf-8-sig"))
    if not found:
        raise ConversionError("Keine Markdown-Tabelle in der Datei gefunden.")
    return found


_NUMBER = re.compile(r"^-?\d+(\.\d+)?$")


def _typed(value: str) -> object:
    if _NUMBER.match(value):
        return float(value) if "." in value else int(value)
    return value


def _sheet_name(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", title).strip()[:31] or "Tabelle"
    name, n = base, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(name.lower())
    return name


@_register("md", "xlsx")
def _md_to_xlsx(src: Path, out: Path, opt: Options) -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for table in _read_tables(src):
        ws = wb.create_sheet(_sheet_name(table.title, used))
        for i, row in enumerate(table.rows):
            ws.append(row if i == 0 else [_typed(c) for c in row])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(
                max(width + 2, 8), 60
            )
    wb.save(out)
    return [out]


@_register("xlsx", "md")
def _xlsx_to_md(src: Path, out: Path, opt: Options) -> list[Path]:
    from openpyxl import load_workbook

    wb = load_workbook(src, read_only=True, data_only=True)
    parts = [f"# {src.stem}"]
    for ws in wb.worksheets:
        rows = tables.trim([list(r) for r in ws.iter_rows(values_only=True)])
        if rows:
            parts.append(f"## {ws.title}\n\n{tables.to_markdown(rows)}")
    wb.close()
    if len(parts) == 1:
        raise ConversionError("Die Arbeitsmappe enthält keine Daten.")
    out.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    return [out]


@_register("md", "csv")
def _md_to_csv(src: Path, out: Path, opt: Options) -> list[Path]:
    found = _read_tables(src)
    # Eine Tabelle -> <name>.csv, mehrere -> <name>_1.csv, <name>_2.csv, ...
    targets = (
        [out]
        if len(found) == 1
        else [
            out.with_name(f"{out.stem}_{i}{out.suffix}")
            for i in range(1, len(found) + 1)
        ]
    )
    for table, target in zip(found, targets):
        _check_target(target, opt)
        with target.open("w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f, delimiter=opt.csv_delimiter).writerows(table.rows)
    return targets


@_register("csv", "md")
def _csv_to_md(src: Path, out: Path, opt: Options) -> list[Path]:
    text = src.read_text(encoding="utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = tables.trim(list(csv.reader(text.splitlines(), dialect)))
    if not rows:
        raise ConversionError("Die CSV-Datei ist leer.")
    out.write_text(f"# {src.stem}\n\n{tables.to_markdown(rows)}\n", encoding="utf-8")
    return [out]


# ---------- PDF -> MD (optional, Docling) ----------


@_register("pdf", "md")
def _pdf_to_md(src: Path, out: Path, opt: Options) -> list[Path]:
    try:
        import importlib

        docling_dc = importlib.import_module("docling.document_converter")
        DocumentConverter = docling_dc.DocumentConverter
    except ImportError:
        raise ConversionError(
            "Der PDF-Import ist in diesem Build nicht enthalten "
            "(Entwicklung: 'uv sync --extra pdf', EXE: mit -WithPdf bauen)."
        ) from None
    result = DocumentConverter().convert(src)
    out.write_text(result.document.export_to_markdown(), encoding="utf-8")
    return [out]


# ---------- Öffentliche API ----------


def _check_target(target: Path, opt: Options) -> None:
    if target.exists() and not opt.overwrite:
        raise ConversionError(
            f"Ziel existiert bereits: {target.name} (Überschreiben aktivieren)"
        )


def convert(
    src: Path, dst_fmt: str, out_dir: Path | None = None, options: Options | None = None
) -> list[Path]:
    """Konvertiert eine Datei und gibt die erzeugten Dateien zurück."""
    opt = options or Options()
    src = Path(src).resolve()
    src_fmt = format_of(src)
    converter = _CONVERTERS.get((src_fmt, dst_fmt))
    if converter is None:
        raise ConversionError(
            f"{src_fmt or src.suffix} → {dst_fmt} wird nicht unterstützt."
        )
    target_dir = Path(out_dir).resolve() if out_dir else src.parent
    target_dir.mkdir(parents=True, exist_ok=True)
    out = target_dir / f"{src.stem}.{dst_fmt}"
    if dst_fmt != "csv":  # CSV prüft pro Teildatei selbst
        _check_target(out, opt)
    try:
        return converter(src, out, opt)
    except (ConversionError, ToolMissingError):
        raise
    except Exception as exc:
        raise ConversionError(str(exc)) from exc


def collect_inputs(
    paths: Iterable[Path], dst_fmt: str, recursive: bool = False
) -> list[Path]:
    """Löst Dateien und Ordner in die konvertierbaren Einzeldateien auf."""
    files: list[Path] = []
    for p in map(Path, paths):
        if p.is_dir():
            pattern = p.rglob("*") if recursive else p.glob("*")
            files += sorted(
                f
                for f in pattern
                if f.is_file() and (format_of(f), dst_fmt) in _CONVERTERS
            )
        else:
            files.append(p)
    return list(dict.fromkeys(files))
